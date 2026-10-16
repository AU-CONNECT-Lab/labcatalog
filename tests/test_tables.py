import io

from docx import Document

from labcatalog.gui.common import col
from labcatalog.tables import PaperTable, csv, docx, files, html, latex, plain, segments

TABLE = PaperTable(
    title="Mood by Group",
    columns=[col("group", "Group"), col("n", "n", numeric=True), col("m", "M (SD)", numeric=True)],
    rows=[
        {"group": "Girls & boys", "n": 61, "m": "8.71 (5.62)"},
        {"group": "  Another_gender", "n": 3},
    ],
    note="*p < .05. α on everyone who answered every item.",
)


def test_statistical_symbols_are_italic_only_as_whole_words():
    assert segments("r [95% CI]") == [("r", True), (" [95% CI]", False)]
    assert segments("Slope (SE)") == [("Slope (", False), ("SE", True), (")", False)]
    assert segments("η² [95% CI]")[0] == ("η²", True)
    assert segments("Parent report") == [("Parent report", False)]


def test_html_draws_three_rules_and_no_vertical_lines():
    text = html(TABLE)
    assert text.startswith('<p style="margin:0"><b>Table X</b></p>')
    assert "<i>Mood by Group</i>" in text
    assert text.count("border-top:1px solid black") == 3
    # Header cells carry the rule above and below, the last row's cells the closing rule.
    assert text.count("border-bottom:1px solid black") == 3 + 3
    assert "border-left" not in text and "border-right" not in text
    assert "Girls &amp; boys" in text
    assert "<i>M</i> (<i>SD</i>)" in text
    assert "<i>Note.</i> *<i>p</i> &lt; .05." in text


def test_page_html_follows_the_text_color_and_drops_the_label():
    text = html(TABLE, page=True)
    assert "Table X" not in text
    assert "currentColor" in text and "black" not in text


def test_plain_text_is_tab_separated():
    assert plain(TABLE).splitlines()[2:5] == [
        "Group\tn\tM (SD)",
        "Girls & boys\t61\t8.71 (5.62)",
        "  Another_gender\t3\t",
    ]


def test_latex_uses_booktabs_and_escapes():
    text = latex(TABLE)
    assert r"\begin{tabular}{lcc}" in text
    assert [line for line in text.splitlines() if line.endswith("rule")] == [
        r"\toprule",
        r"\midrule",
        r"\bottomrule",
    ]
    assert r"Girls \& boys & 61 & 8.71 (5.62) \\" in text
    assert r"~~Another\_gender & 3 &  \\" in text
    assert r"\textit{Note.} *\textit{p} $<$ .05. $\alpha$ on everyone" in text


def test_latex_writes_the_greek_of_fit_tables():
    shown = PaperTable("Fit", [{"name": "d", "label": "Δχ² (Δdf)"}], [{"d": "ω"}])
    text = latex(shown)
    assert r"$\Delta$$\chi$\textsuperscript{2} ($\Delta$df) \\" in text
    assert r"$\omega$ \\" in text


def test_docx_holds_the_label_title_table_and_note():
    document = Document(io.BytesIO(docx(TABLE)))
    texts = [p.text for p in document.paragraphs]
    assert texts[:2] == ["Table X", "Mood by Group"]
    assert texts[-1] == "Note. *p < .05. α on everyone who answered every item."
    [grid] = document.tables
    assert [c.text for c in grid.rows[0].cells] == ["Group", "n", "M (SD)"]
    assert [c.text for c in grid.rows[2].cells] == ["  Another_gender", "3", ""]
    header = grid.rows[0].cells[1]._tc.tcPr.xml
    assert "w:top" in header and "w:bottom" in header
    assert "w:bottom" in grid.rows[2].cells[0]._tc.tcPr.xml
    assert (
        grid.rows[1]
        .cells[0]
        ._tc.tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders")
        is None
    )


def test_csv_uses_the_column_labels_and_opens_in_excel():
    content = csv(TABLE)
    assert content.startswith(b"\xef\xbb\xbf")
    assert content.decode("utf-8-sig").splitlines()[:2] == [
        "Group,n,M (SD)",
        "Girls & boys,61,8.71 (5.62)",
    ]


def test_each_table_is_saved_in_three_formats():
    assert sorted(files("groups", TABLE)) == ["groups.csv", "groups.docx", "groups.tex"]
